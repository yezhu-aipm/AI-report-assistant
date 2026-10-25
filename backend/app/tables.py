"""把 Excel / CSV / Word 表格读成二维列表，再按我们的模板格式解析。不经过大模型，最准也最省钱。"""
import csv, io, re

TRUE = {'✔', '√', '✓', '1', '对', 'T', 't', 'true', 'TRUE', 'Y', 'y', '是', 'o', 'O', '○'}
FALSE = {'×', 'x', 'X', '✗', '✘', '0', '错', 'F', 'f', 'false', 'FALSE', 'N', 'n', '否'}

def cell(v):
    if v is None: return ''
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def read_rows(path, name):
    ext = name.lower().rsplit('.', 1)[-1]
    if ext in ('xlsx', 'xlsm'):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        rows = [[cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
        # 合并单元格：把左上角的值复制到整个区域，方便识别「第1讲」横跨多列的表头
        for mr in ws.merged_cells.ranges:
            v = rows[mr.min_row - 1][mr.min_col - 1] if mr.min_row - 1 < len(rows) else ''
            for r in range(mr.min_row - 1, min(mr.max_row, len(rows))):
                for c in range(mr.min_col - 1, mr.max_col):
                    if c < len(rows[r]) and not rows[r][c]: rows[r][c] = v
        return rows
    if ext == 'xls':
        raise ValueError('暂不支持旧版 .xls，请在 Excel 里另存为 .xlsx 后再上传')
    if ext == 'csv':
        raw = open(path, 'rb').read()
        for enc in ('utf-8-sig', 'gbk'):
            try: text = raw.decode(enc); break
            except UnicodeDecodeError: continue
        return [[cell(v) for v in r] for r in csv.reader(io.StringIO(text))]
    if ext == 'docx':
        import docx
        d = docx.Document(path)
        if not d.tables: return []
        t = max(d.tables, key=lambda t: len(t.rows))
        return [[cell(c.text) for c in r.cells] for r in t.rows]
    return []

def docx_text(path):
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for r in t.rows: parts.append(' | '.join(c.text.strip() for c in r.cells))
    return '\n'.join(parts)

def norm_id(s):
    s = cell(s).replace('（', '(').replace('）', ')').replace(' ', '')
    s = re.sub(r'^第(\d+)题', r'\1', s)
    return s

def qnum(i):
    m = re.match(r'(\d+)', i); return int(m.group(1)) if m else 0

def num(v):
    v = cell(v)
    if v == '': return None
    try: return float(v)
    except ValueError: return None

def _header(rows, keys=('姓名', '学生', '名字')):
    for r, row in enumerate(rows[:15]):
        for c, v in enumerate(row):
            if v.replace(' ', '') in keys: return r, c
    return None, None

def _guess_meta(rows, upto):
    segs = []
    for row in rows[:upto]:
        for v in row:
            if v: segs += [x.strip() for x in re.split(r'[　|,，;；]|\s{2,}', v) if x.strip()]
    meta = {}
    for sg in segs:
        m = re.match(r'^(?:班级|班型)[:：]\s*(.+)$', sg)
        if m and 'cls' not in meta: meta['cls'] = m.group(1).strip()
        m = re.match(r'^(?:考试|测试|考试名称)[:：]\s*(.+)$', sg)
        if m and 'exam' not in meta: meta['exam'] = re.sub(r'[（(].*$', '', m.group(1)).strip()
    if 'cls' not in meta:
        m = next((re.search(r'([A-Za-z]?\d*\s*[\u4e00-\u9fa5]{1,6}班)', sg) for sg in segs if '班' in sg), None)
        if m: meta['cls'] = m.group(1).strip()
    if 'exam' not in meta:
        m = next((sg for sg in segs if any(k in sg for k in ('期中', '期末', '月考', '单元'))), None)
        if m: meta['exam'] = m
    return meta

def parse_grade(rows):
    """逐题成绩表：表头一行「姓名 | 1 | 2 | 14(1) … | 总分」，可选一行「满分」，其下每行一位学生。"""
    hr, nc = _header(rows)
    if hr is None: raise ValueError('成绩表里没找到「姓名」这一列')
    head = rows[hr]
    cols, labels, total_col = [], [], None
    for c in range(nc + 1, len(head)):
        h = head[c].replace(' ', '')
        if not h: continue
        if h in ('总分', '合计', '总计', '得分', '成绩'): total_col = c; continue
        if any(k in h for k in ('排名', '备注', '等级', '学号')): continue
        cols.append(c); labels.append(norm_id(h))
    full, students, scores, official, blanks = None, [], {}, {}, []
    for row in rows[hr + 1:]:
        if nc >= len(row): continue
        name = row[nc].replace(' ', '')
        if not name: continue
        vals = [num(row[c]) if c < len(row) else None for c in cols]
        if '满分' in name or name in ('分值', '满分值'):
            full = vals; continue
        if name in ('平均', '平均分', '班均', '合计'): continue
        students.append(name)
        for k, v in enumerate(vals):
            if v is None: blanks.append({'stu': name, 'k': k})
        scores[name] = [v if v is not None else 0 for v in vals]
        official[name] = num(row[total_col]) if total_col is not None and total_col < len(row) else None
    if not students: raise ValueError('成绩表里没读到学生')
    return dict(labels=labels, full=full, students=students, scores=scores, official=official, blanks=blanks, meta=_guess_meta(rows, hr))

LECT_RE = re.compile(r'(第\s*[\d一二三四五六七八九十]+\s*讲)\s*[:：·\-—\s]*(.*)')

def short_name(name):
    s = re.sub(r'(综合进阶|进阶|专题|提高|拓展|复习)$', '', name).strip()
    return s or name

def parse_exit(rows):
    """出门测记录表：第一行「姓名 | 第1讲 整除综合进阶（合并单元格横跨该讲所有知识点）| …」，
    第二行是每个知识点，下面每行一位学生，✔/× 或 1/0。"""
    hr, nc = _header(rows)
    if hr is None: raise ValueError('出门测表里没找到「姓名」这一列')
    head, kprow = rows[hr], rows[hr + 1] if hr + 1 < len(rows) else []
    lectures, colmap, cur = [], [], None
    for c in range(nc + 1, len(head)):
        h = head[c]
        kp = kprow[c] if c < len(kprow) else ''
        if not h and not kp: continue
        if any(k in (kp or h) for k in ('正确率', '合计', '得分', '总分')): continue
        m = LECT_RE.match(h.replace(' ', '')) if h else None
        key = h or (cur['_h'] if cur else '')
        if cur is None or key != cur['_h']:
            no, nm = (m.group(1).replace(' ', ''), m.group(2).strip()) if m else (f'第{len(lectures) + 1}讲', h)
            cur = {'_h': key, 'no': no, 'name': nm or no, 'kps': []}; lectures.append(cur)
        cur['kps'].append(kp or f'知识点{len(cur["kps"]) + 1}')
        colmap.append((len(lectures) - 1, c))
    for L in lectures: L.pop('_h'); L['short'] = short_name(L['name'])
    students, marks, blanks = [], {}, []
    for row in rows[hr + 2:]:
        if nc >= len(row): continue
        name = row[nc].replace(' ', '')
        if not name or name in ('正确率', '合计', '平均'): continue
        students.append(name)
        mk = [[] for _ in lectures]
        for l, c in colmap:
            v = row[c] if c < len(row) else ''
            if v in TRUE: mk[l].append(1)
            elif v in FALSE: mk[l].append(0)
            else: mk[l].append(0); blanks.append({'stu': name, 'l': l, 'j': len(mk[l]) - 1})
        marks[name] = mk
    if not lectures: raise ValueError('出门测表里没识别出讲次')
    return dict(lectures=lectures, students=students, marks=marks, blanks=blanks)

def parse_paper_table(rows):
    """可选：试卷也可以直接上传题目表「题号 | 题型 | 分值 | 知识点 | 讲次」。"""
    hr = next((r for r, row in enumerate(rows[:10]) if any(v.replace(' ', '') == '题号' for v in row)), None)
    if hr is None: return None
    head = [v.replace(' ', '') for v in rows[hr]]
    def col(*ks): return next((i for i, h in enumerate(head) if any(k in h for k in ks)), None)
    ci, ct, cf, ck, cl = col('题号'), col('题型'), col('分值', '满分'), col('知识点', '考点'), col('讲次', '所属')
    items = []
    for row in rows[hr + 1:]:
        if ci is None or ci >= len(row) or not row[ci]: continue
        i = norm_id(row[ci])
        items.append(dict(id=i, q=qnum(i), type=(row[ct] if ct is not None and ct < len(row) else '') or '解答',
                          full=num(row[cf]) if cf is not None and cf < len(row) else None,
                          kp=row[ck] if ck is not None and ck < len(row) else '', lect_name=row[cl] if cl is not None and cl < len(row) else ''))
    return items or None
