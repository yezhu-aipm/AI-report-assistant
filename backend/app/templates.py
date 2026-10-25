"""给老师下载的 Excel 模板。按模板填写，系统不用大模型就能直接读准。"""
import io
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

HEAD = PatternFill('solid', fgColor='EEF1F6'); BOLD = Font(bold=True)
THIN = Border(*(Side(style='thin', color='D3DAE5'),) * 4)

def _save(wb):
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

def grade_template():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '逐题成绩'
    ws['A1'] = '班级：G6 卓越班　　考试：26秋 期中测试　　（第一行写班级和考试名称，系统会自动识别）'
    labels = ['1', '2', '3', '4', '5', '6', '7', '8', '14(1)', '14(2)']
    head = ['姓名'] + labels + ['总分']
    ws.append([]); ws.append(head); ws.append(['满分', 3, 3, 3, 3, 3, 3, 3, 3, 4, 4, 32])
    ws.append(['张三', 3, 3, 0, 3, 3, 3, 1.5, 3, 4, 2, 25.5]); ws.append(['李四', 3, 0, 3, 3, 3, 3, 3, 3, 4, 4, 29])
    for row in ws.iter_rows(min_row=3, max_row=6):
        for c in row: c.border = THIN; c.alignment = Alignment(horizontal='center')
    for c in ws[3]: c.fill = HEAD; c.font = BOLD
    for c in ws[4]: c.font = Font(color='5F6A82')
    ws['A8'] = '说明：表头写题号（小题写成 14(1)），「满分」一行可选；最后一列「总分」用来核对。'
    ws.column_dimensions['A'].width = 12
    return _save(wb)

def exit_template():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = '出门测'
    lect = [('第1讲 整除综合进阶', ['整除', '因数个数', '整除特征']), ('第2讲 分数综合进阶', ['有限小数化分数', '分数乘除', '分率与具体量'])]
    r1, r2, c = ['姓名'], [''], 2
    for name, kps in lect:
        r1 += [name] + [''] * (len(kps) - 1); r2 += kps
    ws.append(r1); ws.append(r2)
    c = 2
    for name, kps in lect:
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + len(kps) - 1); c += len(kps)
    ws.merge_cells('A1:A2')
    ws.append(['张三', '✔', '✔', '×', '✔', '×', '✔']); ws.append(['李四', '✔', '×', '✔', '✔', '✔', '✔'])
    for row in ws.iter_rows(min_row=1, max_row=4):
        for cell in row: cell.border = THIN; cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in (1, 2):
        for cell in ws[r]: cell.fill = HEAD; cell.font = BOLD
    ws['A6'] = '说明：第一行写讲次（合并单元格横跨这一讲的所有知识点），第二行写知识点；答对填 ✔ 或 1，答错填 × 或 0。'
    for col in 'BCDEFG': ws.column_dimensions[col].width = 14
    return _save(wb)
